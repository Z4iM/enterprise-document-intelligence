"""Structural upload checks; all accepted files remain untrusted parser inputs."""

import json
import os
import signal
import subprocess
import sys
import threading
import zipfile
from io import BytesIO
from pathlib import Path, PureWindowsPath

from fastapi import HTTPException, UploadFile
from starlette.concurrency import run_in_threadpool

from app.ingestion.file_detector import detect_file_type

MAX_FILE = 20 * 1024 * 1024
MAX_PAGES = 20
MAX_PIXELS = 4_000_000
PARSER_SLOTS = threading.BoundedSemaphore(2)


async def read_upload(file: UploadFile) -> tuple[bytes, str, str]:
    name = file.filename or ""
    if (
        not name
        or name in {".", ".."}
        or PureWindowsPath(name).is_reserved()
        or name.endswith((" ", "."))
        or any(ord(c) < 32 or c in '<>"|?*' for c in name)
        or any(c in name for c in ("/", "\\", "\x00", ":"))
        or len(name) > 200
    ):
        raise HTTPException(400, "Invalid filename.")
    kind = detect_file_type(name)
    if kind == "unsupported":
        raise HTTPException(400, "Unsupported file type.")
    data = bytearray()
    while chunk := await file.read(65536):
        data.extend(chunk)
        if len(data) > MAX_FILE:
            raise HTTPException(413, "File exceeds 20 MB limit.")
    if not data:
        raise HTTPException(400, "Empty file.")
    raw = bytes(data)
    try:
        mime = await run_in_threadpool(run_worker, raw, kind, True)
    except Exception as error:
        raise HTTPException(
            400, "Malformed, mismatched or over-budget file."
        ) from error
    return raw, kind, mime


def validate_structure(raw: bytes, kind: str) -> str:
    if kind == "txt":
        text = raw.decode("utf-8")
        if "\x00" in text or any(ord(c) < 32 and c not in "\n\r\t" for c in text):
            raise ValueError("Binary text")
        return "text/plain"
    if kind == "pdf":
        from pypdf import PdfReader

        if not raw.startswith(b"%PDF-"):
            raise ValueError("PDF signature")
        reader = PdfReader(BytesIO(raw), strict=True)
        if reader.is_encrypted or not 0 < len(reader.pages) <= MAX_PAGES:
            raise ValueError("PDF page budget")
        return "application/pdf"
    if kind in {"docx", "pptx"}:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            items = archive.infolist()
            expected = "word/document.xml" if kind == "docx" else "ppt/presentation.xml"
            if expected not in archive.namelist() or len(items) > 1000:
                raise ValueError("Office format")
            if sum(i.file_size for i in items) > MAX_FILE:
                raise ValueError("Archive expansion budget")
            for item in items:
                if (
                    item.flag_bits & 1
                    or ".." in Path(item.filename).parts
                    or item.filename.startswith("/")
                ):
                    raise ValueError("Archive entry")
                if item.file_size > 100 * max(item.compress_size, 1):
                    raise ValueError("Compression ratio")
        if kind == "docx":
            from docx import Document

            Document(BytesIO(raw))
        else:
            from pptx import Presentation

            if len(Presentation(BytesIO(raw)).slides) > MAX_PAGES:
                raise ValueError("Slide budget")
        return "application/vnd.openxmlformats-officedocument." + (
            "wordprocessingml.document"
            if kind == "docx"
            else "presentationml.presentation"
        )
    from PIL import Image

    with Image.open(BytesIO(raw)) as image:
        if (
            image.format not in {"PNG", "JPEG"}
            or image.width * image.height > MAX_PIXELS
        ):
            raise ValueError("Image budget")
        image.verify()
        return "image/png" if image.format == "PNG" else "image/jpeg"


def run_worker(raw: bytes, kind: str, validate_only: bool = False):
    if not PARSER_SLOTS.acquire(timeout=2):
        raise ValueError("Parser capacity exhausted.")
    env = {
        key: os.environ[key]
        for key in ("PATH", "SYSTEMROOT", "WINDIR", "TEMP", "TMP")
        if key in os.environ
    }
    env.update(
        PYTHONPATH=str(Path(__file__).resolve().parents[2]),
        PYTHON_DOTENV_DISABLED="1",
        PYTHONUTF8="1",
    )
    options = (
        # Language names only; no application credentials enter parser children.
        {"creationflags": subprocess.CREATE_NO_WINDOW}
        if os.name == "nt"
        else {"start_new_session": True}
    )
    env["EDI_OCR_LANGUAGES"] = os.getenv("EDI_OCR_LANGUAGES", "eng+ara")
    try:
        with subprocess.Popen(
            [
                sys.executable,
                "-m",
                "app.ingestion.worker",
                kind,
                "validate" if validate_only else "extract",
            ],
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.DEVNULL,
            env=env,
            **options,
        ) as process:
            try:
                output, _ = process.communicate(raw, timeout=45)
            except subprocess.TimeoutExpired:
                if os.name == "nt":
                    subprocess.run(
                        ["taskkill", "/PID", str(process.pid), "/T", "/F"],
                        stdout=subprocess.DEVNULL,
                        stderr=subprocess.DEVNULL,
                        check=False,
                    )
                else:
                    os.killpg(process.pid, signal.SIGKILL)
                process.kill()
                process.wait()
                raise ValueError("Parser deadline exceeded.") from None
            if process.returncode:
                raise ValueError("Parser rejected input.")
            return json.loads(output)
    finally:
        PARSER_SLOTS.release()


def extract_text_isolated(file_bytes: bytes, file_type: str) -> tuple[str, str]:
    """Disposable, bounded parser process; no credentials or application database."""
    text, method = run_worker(file_bytes, file_type)
    return text, method
